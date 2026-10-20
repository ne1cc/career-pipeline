"""Job-listing enrichment for the daily pipeline (apply type, new jobs, fresh posts)."""

from __future__ import annotations

import json
import re
from datetime import date, datetime, timedelta, timezone
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl.styles import PatternFill

NEW_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FRESH_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
NEW_AND_FRESH_FILL = PatternFill(start_color="B4D7FF", end_color="B4D7FF", fill_type="solid")

ENRICHMENT_COLUMNS = ("apply_type", "is_new", "is_fresh_post")
FRESH_POST_DAYS = 2
ANALYST_PATTERN = r"\b(?:analyst|analytics|data)\b"


def load_previous_urls(state_path: Path) -> dict[str, set[str]]:
    if not state_path.exists():
        return {}
    try:
        payload = json.loads(state_path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {}
    segments = payload.get("segments", {})
    return {name: set(urls) for name, urls in segments.items() if isinstance(urls, list)}


def save_current_urls(state_path: Path, segment_urls: dict[str, Iterable[str]]) -> None:
    state_path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "last_run": datetime.now().isoformat(timespec="seconds"),
        "segments": {name: sorted(set(urls)) for name, urls in segment_urls.items()},
    }
    state_path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def parse_posted_date(value) -> date | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    if isinstance(value, datetime):
        return value.date()
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y"):
        try:
            return datetime.strptime(text[:19], fmt).date()
        except ValueError:
            continue
    return None


def classify_apply_type(row: pd.Series) -> str:
    site = str(row.get("site", "") or "").lower()
    job_url = str(row.get("job_url", "") or "").lower()
    job_url_direct = str(row.get("job_url_direct", "") or "").strip().lower()
    combined_urls = f"{job_url} {job_url_direct}"

    if "greenhouse.io" in combined_urls:
        return "Greenhouse Direct"
    if "lever.co" in combined_urls:
        return "Lever Direct"
    if "ashbyhq.com" in combined_urls:
        return "Ashby Direct"
    if "myworkdayjobs.com" in combined_urls or "workday" in combined_urls:
        return "Workday Direct"
    if "smartrecruiters.com" in combined_urls:
        return "SmartRecruiters Direct"
    if "builtin" in combined_urls or "builtin" in site:
        return "BuiltIn Direct"
    if "himalayas" in combined_urls or "himalayas" in site:
        return "Himalayas Direct"
    if "remotive" in combined_urls or "remotive" in site:
        return "Remotive Direct"
    if "jobicy" in combined_urls or "jobicy" in site:
        return "Jobicy Direct"
    if "remoteok" in combined_urls or "remoteok" in site:
        return "RemoteOK Direct"
    if "arbeitnow" in combined_urls or "arbeitnow" in site:
        return "Arbeitnow Direct"
    if "hiringcafe" in combined_urls or "hiringcafe" in site:
        return "HiringCafe Direct"

    if site == "ats":
        return "Company ATS Direct"

    if site == "indeed":
        if "indeed-apply" in job_url_direct or "indeed-apply" in job_url:
            return "Indeed Easy Apply"
        if job_url_direct and job_url_direct not in ("", "nan", "none"):
            return "Company Site"
        return "Indeed (unknown apply)"

    if site == "linkedin":
        if not job_url_direct or job_url_direct in ("nan", "none"):
            return "LinkedIn Easy Apply"
        return "Company Site"

    if site == "zip_recruiter":
        return "ZipRecruiter Apply"

    if site == "glassdoor":
        return "Glassdoor Apply"

    if site == "google":
        return "Google Jobs (redirect)"

    if job_url_direct and job_url_direct not in ("", "nan", "none"):
        return "Company Site"
    return "Unknown"


def enrich_dataframe(
    df: pd.DataFrame,
    previous_urls: set[str],
    *,
    reference_date: date | None = None,
    has_previous_run: bool,
) -> pd.DataFrame:
    if df.empty:
        for col in ENRICHMENT_COLUMNS:
            df[col] = pd.Series(dtype=str if col == "apply_type" else bool)
        return df

    today = reference_date or date.today()
    fresh_cutoff = today - timedelta(days=FRESH_POST_DAYS)

    enriched = df.copy()
    enriched["apply_type"] = enriched.apply(classify_apply_type, axis=1)

    if "job_url" in enriched.columns and has_previous_run:
        enriched["is_new"] = ~enriched["job_url"].astype(str).isin(previous_urls)
    else:
        enriched["is_new"] = False

    posted_dates = enriched.get("date_posted", pd.Series(index=enriched.index)).map(parse_posted_date)
    enriched["is_fresh_post"] = posted_dates.map(
        lambda posted: bool(posted and posted >= fresh_cutoff)
    )

    front = [col for col in ENRICHMENT_COLUMNS if col in enriched.columns]
    other = [col for col in enriched.columns if col not in front]
    return enriched[front + other]


TRACK_A_PATTERN = r"\b(?:analyst|analytics|data|data scientist|statistician|quantitative|analytics engineer|bi\b|bi developer|business intelligence|reporting|insights)\b"
TRACK_C_PATTERN = r"\b(?:product manager|apm\b|business operations|bizops|operations analyst|project coordinator|strategy & operations|strategy and operations|strategy analyst)\b"
TRACK_B_PATTERN = r"\b(?:software engineer|developer|swe\b|frontend|backend|full stack)\b"

DEVELOPER_PATTERN = (
    r"\b(?:"
    r"software engineer|software developer|software dev\b|swe\b|sde\b|"
    r"developer|programmer|"
    r"frontend|front end|front-end|"
    r"backend|back end|back-end|"
    r"full stack|fullstack|full-stack|"
    r"web developer|mobile developer|ios developer|android developer|"
    r"cloud engineer|cloud developer|devops|sre\b|site reliability|"
    r"platform engineer|infrastructure engineer|systems engineer|systems software|"
    r"application developer|application engineer|apps developer|"
    r"firmware engineer|embedded software|embedded engineer|"
    r"machine learning engineer|ml engineer|ai engineer"
    r")\b"
)
SECURITY_PATTERN = (
    r"\b(?:"
    r"security|cybersecurity|cyber security|infosec|information security|"
    r"application security|appsec|cloud security|network security|"
    r"security engineer|security analyst|soc analyst|security operations|"
    r"penetration tester|pentester|vulnerability|threat intelligence|"
    r"incident response|cyber defense|product security"
    r")\b"
)
DEV_EXCLUSION_PATTERN = (
    r"\b(?:"
    r"business development|sales development|real estate development|"
    r"donor development|fund development|biz dev|sdr\b|bdr\b"
    r")\b"
)

SENIOR_TITLE_PATTERN = r"\b(?:senior|sr\b|sr\.|lead|principal|staff|(?<!product\s)(?<!project\s)manager|director|vp\b|vice president|head of|chief|architect|distinguished|fellow|iii|iv|v)\b"
HIGH_YOE_PATTERN = r"\b(?:[5-9]|1[0-9])\+\s*(?:years|yrs)\b|\b(?:[5-9]|1[0-9])\s*(?:to|-|–)\s*[0-9]+\s*(?:years|yrs)\b"
ENTRY_TITLE_TAG_PATTERN = r"\b(?:new grad|junior|associate|entry level|early career|university|intern|internship|campus)\b"


def matches_fresher_track(title: str) -> bool:
    if not title or pd.isna(title):
        return False
    t = str(title).lower()
    return bool(
        re.search(TRACK_A_PATTERN, t)
        or re.search(TRACK_C_PATTERN, t)
        or re.search(TRACK_B_PATTERN, t)
        or re.search(SECURITY_PATTERN, t)
    )


def matches_developer_security_track(title: str, description: str = "") -> bool:
    if not title or pd.isna(title):
        return False
    t = str(title).lower()
    if re.search(DEV_EXCLUSION_PATTERN, t):
        return False
    if not (re.search(DEVELOPER_PATTERN, t) or re.search(SECURITY_PATTERN, t)):
        return False
    return not is_senior_or_high_yoe(t, description)


def is_senior_or_high_yoe(title: str, description: str = "") -> bool:
    t = "" if pd.isna(title) else str(title or "").lower()
    if re.search(SENIOR_TITLE_PATTERN, t):
        return True
    d = "" if pd.isna(description) else str(description or "").lower()
    if d and re.search(HIGH_YOE_PATTERN, d):
        return True
    return False



def calculate_fresher_priority_score(row: pd.Series) -> int:
    score = 0
    is_new = row.get("is_new", False)
    if pd.notna(is_new) and bool(is_new):
        score += 5
    is_fresh = row.get("is_fresh_post", False)
    if pd.notna(is_fresh) and bool(is_fresh):
        score += 3

    title = "" if pd.isna(row.get("title")) else str(row.get("title", "")).lower()
    if re.search(ENTRY_TITLE_TAG_PATTERN, title):
        score += 3

    apply_type = "" if pd.isna(row.get("apply_type")) else str(row.get("apply_type", "")).lower()
    if any(k in apply_type for k in ["direct", "easy apply"]):
        score += 2

    if re.search(TRACK_A_PATTERN, title) or re.search(TRACK_C_PATTERN, title) or re.search(TRACK_B_PATTERN, title) or re.search(SECURITY_PATTERN, title):
        score += 1

    return score


def calculate_developer_security_priority_score(row: pd.Series) -> int:
    score = 0
    is_new = row.get("is_new", False)
    if pd.notna(is_new) and bool(is_new):
        score += 5
    is_fresh = row.get("is_fresh_post", False)
    if pd.notna(is_fresh) and bool(is_fresh):
        score += 3

    title = "" if pd.isna(row.get("title")) else str(row.get("title", "")).lower()
    if re.search(ENTRY_TITLE_TAG_PATTERN, title):
        score += 4
    elif re.search(r"\b(?:i\b|1\b|associate)\b", title):
        score += 2

    apply_type = "" if pd.isna(row.get("apply_type")) else str(row.get("apply_type", "")).lower()
    if any(k in apply_type for k in ["direct", "easy apply"]):
        score += 2

    if re.search(SECURITY_PATTERN, title):
        score += 2
    if re.search(DEVELOPER_PATTERN, title):
        score += 1

    return score


def deduplicate_action_items(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return df
    df = df.copy()
    # Prefer Direct ATS and higher priority scores
    if "_score" in df.columns:
        df = df.sort_values(by=["_score"], ascending=False)
    # Deduplicate on normalized title and company
    if "title" in df.columns and "company" in df.columns:
        clean_title = df["title"].astype(str).str.lower().str.replace(r"[^\w\s]", "", regex=True)
        clean_company = df["company"].astype(str).str.lower().str.replace(r"[^\w\s]", "", regex=True)
        df["_dedup_key"] = clean_title + "::" + clean_company
        df = df.drop_duplicates(subset=["_dedup_key"], keep="first").drop(columns=["_dedup_key"])
    return df


def analyst_mask(df: pd.DataFrame) -> pd.Series:
    if "title" not in df.columns:
        return pd.Series(False, index=df.index)
    return df["title"].astype(str).str.contains(ANALYST_PATTERN, flags=re.IGNORECASE, regex=True, na=False)


def build_action_items_sheet(segment_frames: dict[str, pd.DataFrame]) -> pd.DataFrame:
    pieces = []
    for segment_name, df in segment_frames.items():
        if df.empty:
            continue
        desc_col = "description" if "description" in df.columns else None

        # Filter for fresher track & exclude senior using boolean mask
        title_series = df["title"].fillna("").astype(str) if "title" in df.columns else pd.Series("", index=df.index)
        desc_series = df[desc_col].fillna("").astype(str) if desc_col else pd.Series("", index=df.index)
        mask = [
            matches_fresher_track(t) and not is_senior_or_high_yoe(t, d)
            for t, d in zip(title_series, desc_series)
        ]

        subset = df[mask].copy()
        if subset.empty:
            continue
        subset.insert(0, "segment", segment_name)
        pieces.append(subset)

    if not pieces:
        return pd.DataFrame(columns=["segment", *ENRICHMENT_COLUMNS])

    combined = pd.concat(pieces, ignore_index=True)
    combined["_score"] = combined.apply(calculate_fresher_priority_score, axis=1)
    combined = deduplicate_action_items(combined)

    if "date_posted" in combined.columns:
        combined["_date_sort"] = pd.to_datetime(combined["date_posted"], errors="coerce")
        combined = combined.sort_values(by=["_score", "_date_sort"], ascending=[False, False]).drop(columns=["_score", "_date_sort"])
    else:
        combined = combined.sort_values(by="_score", ascending=False).drop(columns="_score")

    return combined.reset_index(drop=True)


def build_developer_security_sheet(segment_frames: dict[str, pd.DataFrame]) -> pd.DataFrame:
    pieces = []
    for segment_name, df in segment_frames.items():
        if df.empty:
            continue
        desc_col = "description" if "description" in df.columns else None

        title_series = df["title"].fillna("").astype(str) if "title" in df.columns else pd.Series("", index=df.index)
        desc_series = df[desc_col].fillna("").astype(str) if desc_col else pd.Series("", index=df.index)
        mask = [
            matches_developer_security_track(t, d)
            for t, d in zip(title_series, desc_series)
        ]

        subset = df[mask].copy()
        if subset.empty:
            continue
        subset.insert(0, "segment", segment_name)
        pieces.append(subset)

    if not pieces:
        return pd.DataFrame(columns=["segment", *ENRICHMENT_COLUMNS])

    combined = pd.concat(pieces, ignore_index=True)
    combined["_score"] = combined.apply(calculate_developer_security_priority_score, axis=1)
    combined = deduplicate_action_items(combined)

    if "date_posted" in combined.columns:
        combined["_date_sort"] = pd.to_datetime(combined["date_posted"], errors="coerce")
        combined = combined.sort_values(by=["_score", "_date_sort"], ascending=[False, False]).drop(columns=["_score", "_date_sort"])
    else:
        combined = combined.sort_values(by="_score", ascending=False).drop(columns="_score")

    return combined.reset_index(drop=True)


LAST_72_HOURS_COLUMNS = [
    "title", "company", "location", "posted_at", "date_precision",
    "experience_fit", "job_url", "job_url_direct", "sources", "site",
]


def _parse_fresh_timestamp(value) -> tuple[datetime | None, str]:
    if value is None or pd.isna(value) or not str(value).strip():
        return None, "missing"
    text = str(value).strip()
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        return None, "invalid"
    if parsed.tzinfo is None:
        return parsed.replace(tzinfo=timezone.utc), "timestamp (timezone unknown)"
    return parsed.astimezone(timezone.utc), "timestamp"


def _canonical_job_url(value) -> str:
    text = "" if pd.isna(value) else str(value).strip()
    if not text:
        return ""
    parts = urlsplit(text)
    query = [(key, val) for key, val in parse_qsl(parts.query, keep_blank_values=True)
             if not key.lower().startswith("utm_") and key.lower() not in {"source", "ref", "gh_src"}]
    return urlunsplit((parts.scheme.lower(), parts.netloc.lower(), parts.path.rstrip("/"), urlencode(query), ""))


def build_last_72_hours_sheet(
    segment_frames: dict[str, pd.DataFrame], *, reference_now: datetime | None = None
) -> pd.DataFrame:
    """Build an entry-level view of supportably fresh postings from all sources."""
    now = reference_now or datetime.now(timezone.utc)
    if now.tzinfo is None:
        now = now.replace(tzinfo=timezone.utc)
    now = now.astimezone(timezone.utc)
    cutoff = now - timedelta(hours=72)

    pieces = []
    for segment_name, frame in segment_frames.items():
        if frame.empty or "title" not in frame.columns:
            continue
        df = frame.copy()
        descriptions = df.get("description", pd.Series("", index=df.index)).fillna("").astype(str)
        titles = df["title"].fillna("").astype(str)
        eligible = [
            (matches_fresher_track(title) or matches_developer_security_track(title))
            and not is_senior_or_high_yoe(title, description)
            for title, description in zip(titles, descriptions)
        ]
        df = df[eligible].copy()
        if df.empty:
            continue

        posted_values = df.get("posted_at", pd.Series("", index=df.index))
        date_values = df.get("date_posted", pd.Series("", index=df.index))
        parsed_values = []
        precision_values = []
        keep = []
        for posted_at, posted_date in zip(posted_values, date_values):
            parsed, precision = _parse_fresh_timestamp(posted_at)
            if parsed is None and precision == "missing":
                calendar_date = parse_posted_date(posted_date)
                if calendar_date:
                    parsed = datetime.combine(calendar_date, datetime.min.time(), tzinfo=timezone.utc)
                    precision = "date only (UTC; full-day conservative)"
            parsed_values.append(parsed)
            precision_values.append(precision)
            if not parsed:
                keep.append(False)
            elif precision == "timestamp (timezone unknown)":
                keep.append(cutoff <= parsed - timedelta(hours=14) and parsed + timedelta(hours=14) <= now)
            elif precision.startswith("date only"):
                keep.append(cutoff <= parsed - timedelta(hours=14) and parsed + timedelta(hours=38) <= now)
            else:
                keep.append(cutoff <= parsed <= now)
        df = df[keep].copy()
        if df.empty:
            continue
        kept_positions = [position for position, is_kept in enumerate(keep) if is_kept]
        df["posted_at"] = [parsed_values[i].isoformat() for i in kept_positions]
        df["date_precision"] = [precision_values[i] for i in kept_positions]
        experience_labels = []
        for desc in df.get("description", pd.Series("", index=df.index)).fillna("").astype(str):
            if re.search(r"\b(?:[3-9]|1[0-9])\s*(?:\+|to|-|–)\s*(?:[3-9]|1[0-9])?\s*(?:years|yrs)\b", desc, re.I):
                experience_labels.append("Stretch (3+ years stated)")
            elif re.search(r"\b(?:\d+\+?\s*(?:to|-|–)\s*)?\d+\+?\s*(?:years|yrs)\b", desc, re.I):
                experience_labels.append("Experience stated; review requirements")
            else:
                experience_labels.append("Experience not stated")
        df["experience_fit"] = experience_labels
        df["sources"] = df.get("site", pd.Series(segment_name, index=df.index)).fillna("").astype(str)
        pieces.append(df)

    if not pieces:
        return pd.DataFrame(columns=LAST_72_HOURS_COLUMNS)

    combined = pd.concat(pieces, ignore_index=True, sort=False)
    combined["_dedup_key"] = combined.apply(
        lambda row: _canonical_job_url(row.get("job_url_direct"))
        or _canonical_job_url(row.get("job_url")), axis=1
    )
    combined = combined[combined["_dedup_key"] != ""].copy()
    rows = []
    for _, group in combined.groupby("_dedup_key", sort=False):
        row = group.sort_values("posted_at", ascending=False).iloc[0].copy()
        sources = sorted({source for value in group["sources"] for source in str(value).split(", ") if source})
        row["sources"] = ", ".join(sources)
        rows.append(row)
    result = pd.DataFrame(rows)
    result = result.sort_values("posted_at", ascending=False)
    columns = [col for col in LAST_72_HOURS_COLUMNS if col in result.columns]
    return result[columns].reset_index(drop=True)


def apply_row_highlights(ws, df: pd.DataFrame) -> None:
    """Highlight full rows for new listings and recently posted jobs."""
    if df.empty or ws.max_row < 2:
        return

    col_index = {name: idx + 1 for idx, name in enumerate(df.columns)}
    new_idx = col_index.get("is_new")
    fresh_idx = col_index.get("is_fresh_post")
    if not new_idx or not fresh_idx:
        return

    for row_idx in range(2, ws.max_row + 1):
        is_new = bool(ws.cell(row=row_idx, column=new_idx).value)
        is_fresh = bool(ws.cell(row=row_idx, column=fresh_idx).value)
        if not is_new and not is_fresh:
            continue
        fill = NEW_AND_FRESH_FILL if is_new and is_fresh else NEW_FILL if is_new else FRESH_FILL
        for col_idx in range(1, ws.max_column + 1):
            ws.cell(row=row_idx, column=col_idx).fill = fill


def summarize_enrichment(segment_frames: dict[str, pd.DataFrame]) -> dict[str, int]:
    totals = {
        "total_jobs": 0,
        "new_jobs": 0,
        "fresh_posts": 0,
        "easy_apply": 0,
        "direct_ats": 0,
        "action_items": 0,
        "developer_security": 0,
    }
    for df in segment_frames.values():
        if df.empty:
            continue
        totals["total_jobs"] += len(df)
        if "is_new" in df.columns:
            totals["new_jobs"] += int(df["is_new"].sum())
        if "is_fresh_post" in df.columns:
            totals["fresh_posts"] += int(df["is_fresh_post"].sum())
        if "apply_type" in df.columns:
            totals["easy_apply"] += int(
                df["apply_type"].str.contains("Easy Apply", case=False, na=False).sum()
            )
            totals["direct_ats"] += int(
                df["apply_type"].str.contains("Direct", case=False, na=False).sum()
            )
    action_items = build_action_items_sheet(segment_frames)
    totals["action_items"] = len(action_items)
    dev_sec_items = build_developer_security_sheet(segment_frames)
    totals["developer_security"] = len(dev_sec_items)
    return totals


def format_email_summary(stats: dict[str, int], *, has_previous_run: bool) -> str:
    new_line = (
        f"- New since last run: {stats.get('new_jobs', 0)}\n"
        if has_previous_run
        else "- New since last run: (first run — all jobs treated as baseline)\n"
    )
    direct_ats_count = stats.get("direct_ats", 0)
    dev_sec_count = stats.get("developer_security", 0)
    return (
        f"Quick stats:\n"
        f"- Total Raw Postings: {stats.get('total_jobs', 0)}\n"
        f"- Action Items (Fresher Spotlight): {stats.get('action_items', 0)}\n"
        f"- Developer & Security Roles: {dev_sec_count}\n"
        f"- Direct ATS listings: {direct_ats_count}\n"
        f"{new_line}"
        f"- Posted in last {FRESH_POST_DAYS} days: {stats.get('fresh_posts', 0)}\n"
        f"- Easy Apply listings: {stats.get('easy_apply', 0)}\n\n"
        f"Legend in Excel:\n"
        f"- Green row = new since your last daily run\n"
        f"- Yellow row = posted within {FRESH_POST_DAYS} days\n"
        f"- Blue row = both new and fresh\n"
        f"- apply_type column = Indeed/LinkedIn Easy Apply vs company site / Direct ATS\n"
    )
