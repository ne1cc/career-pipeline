from pathlib import Path
from datetime import datetime, timezone
import pandas as pd
import openpyxl
from run_pipeline import build_excel_workbook, EXPECTED_OUTPUTS
from pipeline_enrichment import format_email_summary, summarize_enrichment


def test_build_excel_workbook_end_to_end(tmp_path):
    results_dir = tmp_path / "results"
    results_dir.mkdir()

    # Create sample CSVs for all 8 expected outputs
    posted_at = datetime.now(timezone.utc).isoformat()
    posted_date = posted_at[:10]
    sample_jobs = [
        {"title": "Junior Data Analyst", "company": "Co A", "job_url": "https://a.com/job1", "date_posted": "2026-09-10", "site": "linkedin"},
        {"title": "Senior Data Scientist", "company": "Co B", "job_url": "https://b.com/job2", "date_posted": "2026-09-08", "site": "indeed"},
        {"title": "Analytics Engineer", "company": "Co C", "job_url": "https://c.com/job3", "date_posted": "2026-09-11", "site": "ats"},
        {"title": "Junior Software Engineer", "company": "Co D", "job_url": "https://d.com/job4", "date_posted": "2026-09-11", "site": "ats"},
        {"title": "Cybersecurity Analyst", "company": "Co E", "job_url": "https://e.com/job5", "date_posted": "2026-09-11", "site": "ats"},
        {"title": "Senior Security Engineer", "company": "Co F", "job_url": "https://f.com/job6", "date_posted": "2026-09-11", "site": "indeed"},
    ]
    df = pd.DataFrame(sample_jobs)
    df["posted_at"] = posted_at
    df.loc[df["title"].str.startswith("Senior"), "posted_at"] = "2026-01-01T00:00:00+00:00"
    df.loc[df["title"].str.startswith("Senior"), "date_posted"] = "2026-01-01"
    df.loc[~df["title"].str.startswith("Senior"), "date_posted"] = posted_date
    for csv_name in EXPECTED_OUTPUTS.keys():
        df.to_csv(results_dir / csv_name, index=False)

    state_file = tmp_path / "state.json"
    excel_path, stats, has_previous_run = build_excel_workbook(results_dir, EXPECTED_OUTPUTS, state_file=state_file)
    assert excel_path.exists()
    assert stats["total_jobs"] == len(EXPECTED_OUTPUTS) * len(sample_jobs)
    # Action Items: Junior Data Analyst, Analytics Engineer, Junior Software Engineer, Cybersecurity Analyst (analyst match)
    assert stats["action_items"] == 4
    # Developer & Security: Junior Software Engineer, Cybersecurity Analyst
    assert stats["developer_security"] == 2
    assert stats["direct_ats"] == 3 * len(EXPECTED_OUTPUTS)

    wb = openpyxl.load_workbook(excel_path)
    sheet_names = wb.sheetnames

    # Verify the fresh sheet leads, followed by the existing priority views.
    assert sheet_names[:3] == ["Last 72 Hours", "Action Items", "Developer & Security"]
    ws_fresh = wb["Last 72 Hours"]
    fresh_header = [cell.value for cell in ws_fresh[1]]
    fresh_title_idx = fresh_header.index("title")
    fresh_titles = [row[fresh_title_idx] for row in ws_fresh.iter_rows(min_row=2, values_only=True)]
    assert "Junior Data Analyst" in fresh_titles
    assert "Senior Data Scientist" not in fresh_titles

    # Verify all raw tabs exist
    for target_tab in EXPECTED_OUTPUTS.values():
        assert target_tab in sheet_names

    # Verify Action Items excluded the Senior role
    ws_action = wb["Action Items"]
    header = [cell.value for cell in ws_action[1]]
    title_idx = header.index("title")
    all_action_titles = [row[title_idx] for row in ws_action.iter_rows(min_row=2, values_only=True)]
    assert "Senior Data Scientist" not in all_action_titles
    assert "Analytics Engineer" in all_action_titles

    # Verify Developer & Security tab contents
    ws_dev_sec = wb["Developer & Security"]
    dev_header = [cell.value for cell in ws_dev_sec[1]]
    dev_title_idx = dev_header.index("title")
    all_dev_titles = [row[dev_title_idx] for row in ws_dev_sec.iter_rows(min_row=2, values_only=True)]
    assert "Junior Software Engineer" in all_dev_titles
    assert "Cybersecurity Analyst" in all_dev_titles
    assert "Senior Security Engineer" not in all_dev_titles
    assert "Junior Data Analyst" not in all_dev_titles


def test_email_summary_highlights():
    sample_df = pd.DataFrame([
        {"title": "Junior Data Analyst", "company": "Co A", "job_url": "https://a.com/job1", "date_posted": "2026-09-10", "site": "ats", "apply_type": "Company ATS Direct", "is_new": True, "is_fresh_post": True},
        {"title": "Analytics Engineer", "company": "Co B", "job_url": "https://b.com/job2", "date_posted": "2026-09-11", "site": "linkedin", "apply_type": "LinkedIn Easy Apply", "is_new": True, "is_fresh_post": True},
        {"title": "Junior Software Engineer", "company": "Co C", "job_url": "https://c.com/job3", "date_posted": "2026-09-11", "site": "ats", "apply_type": "Company ATS Direct", "is_new": True, "is_fresh_post": True},
    ])
    segment_frames = {"Test Segment": sample_df}
    stats = summarize_enrichment(segment_frames)
    summary = format_email_summary(stats, has_previous_run=True)

    assert "Total Raw Postings" in summary
    assert "Action Items (Fresher Spotlight)" in summary
    assert "Developer & Security Roles" in summary
    assert "Direct ATS" in summary
