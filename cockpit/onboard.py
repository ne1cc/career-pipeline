from __future__ import annotations

import re
import shutil
import subprocess
from pathlib import Path
from typing import Any

from cockpit import intel
from cockpit.exporters import build_base_resumes
from cockpit.normalizer import REPO_ROOT


def seed_user_layer(
    repo_root: Path = REPO_ROOT,
    overwrite: bool = False,
) -> list[Path]:
    """Seeds gitignored user-layer files from committed system templates when missing.

    Contract:
    - Copies starter templates into `config/evidence/` and `config/career_ops/` only if missing (unless `overwrite=True`).
    - Returns list of created/updated file paths.
    """
    created: list[Path] = []
    co_root = repo_root / "config" / "career_ops"
    ev_src = repo_root / "config" / "templates" / "evidence"
    ev_dst = repo_root / "config" / "evidence"

    # 1. Seed career_ops user layer from .example / .template files
    seed_pairs = [
        (co_root / "cv.example.md", co_root / "cv.md"),
        (co_root / "config" / "profile.example.yml", co_root / "config" / "profile.yml"),
        (co_root / "modes" / "_profile.template.md", co_root / "modes" / "_profile.md"),
        (co_root / "modes" / "_custom.template.md", co_root / "modes" / "_custom.md"),
        (co_root / "voice-dna.template.md", co_root / "voice-dna.md"),
        (repo_root / "pipeline_config.example.json", repo_root / "pipeline_config.json"),
    ]
    for src, dst in seed_pairs:
        if src.exists() and (overwrite or not dst.exists()):
            dst.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, dst)
            created.append(dst)

    # 2. Seed config/evidence/ from config/templates/evidence/
    if ev_src.exists():
        for src_file in sorted(ev_src.rglob("*")):
            if not src_file.is_file():
                continue
            rel = src_file.relative_to(ev_src)
            dst_file = ev_dst / rel
            if overwrite or not dst_file.exists():
                dst_file.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(src_file, dst_file)
                created.append(dst_file)

    return created


def extract_resume_text(path: Path) -> str:
    """Extracts plain/markdown text from a .md, .txt, .docx, or .pdf resume file.

    Contract:
    - Returns extracted UTF-8 text string.
    - Raises ValueError on unsupported file extensions or unreadable files.
    """
    p = Path(path).expanduser().resolve()
    if not p.exists():
        raise FileNotFoundError(f"Resume file not found: {p}")

    ext = p.suffix.lower()
    if ext in (".md", ".txt", ".tex"):
        return p.read_text("utf-8", errors="replace").strip()

    if ext == ".docx":
        import docx

        doc = docx.Document(str(p))
        lines: list[str] = []
        for para in doc.paragraphs:
            txt = para.text.strip()
            if txt:
                lines.append(txt)
        for table in doc.tables:
            for row in table.rows:
                row_txt = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if row_txt:
                    lines.append(row_txt)
        return "\n".join(lines).strip()

    if ext == ".pdf":
        # 1. Try pypdf first
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(p))
            pages = [page.extract_text() or "" for page in reader.pages]
            joined = "\n".join(t.strip() for t in pages if t.strip()).strip()
            if joined:
                return joined
        except Exception:
            pass

        # 2. Try pdftotext fallback if installed
        try:
            proc = subprocess.run(
                ["pdftotext", "-layout", str(p), "-"],
                capture_output=True,
                text=True,
                timeout=20,
                check=False,
            )
            if proc.returncode == 0 and proc.stdout.strip():
                return proc.stdout.strip()
        except (FileNotFoundError, OSError, subprocess.SubprocessError):
            pass

        raise ValueError(f"Could not extract text from PDF: {p}. Install pypdf or poppler (pdftotext).")

    raise ValueError(f"Unsupported resume format '{ext}'. Use .pdf, .docx, .md, or .txt.")


def parse_contact_heuristics(text: str) -> dict[str, str]:
    """Heuristically extracts candidate name, email, phone, github, and linkedin from resume text."""
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    name = "Candidate"
    if lines:
        first = re.sub(r"^[#*_`\s]+|[#*_`\s]+$", "", lines[0]).strip()
        if 1 <= len(first.split()) <= 5 and "@" not in first and "http" not in first.lower():
            name = first

    email_m = re.search(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", text)
    phone_m = re.search(r"(\+?\d[\d\s().-]{7,}\d)", text)
    gh_m = re.search(r"((?:https?://)?(?:www\.)?github\.com/[A-Za-z0-9_.-]+)", text, re.IGNORECASE)
    li_m = re.search(r"((?:https?://)?(?:www\.)?linkedin\.com/in/[A-Za-z0-9_.-]+)", text, re.IGNORECASE)

    return {
        "name": name,
        "email": email_m.group(0) if email_m else "",
        "phone": phone_m.group(1).strip() if phone_m else "",
        "github": gh_m.group(1) if gh_m else "",
        "linkedin": li_m.group(1) if li_m else "",
    }


def write_user_profile_configs(
    *,
    repo_root: Path = REPO_ROOT,
    name: str,
    email: str = "",
    phone: str = "",
    location: str = "Remote",
    linkedin: str = "",
    github: str = "",
    portfolio: str = "",
    primary_roles: list[str] | None = None,
    primary_locations: list[str] | None = None,
    min_comp: str = "$100K",
    target_comp: str = "$120K-160K",
) -> dict[str, Path]:
    """Writes personalized profile.yml, 01_master_career_history.yaml header, and scoring_rules.local.toml."""
    seed_user_layer(repo_root=repo_root, overwrite=False)
    co_root = repo_root / "config" / "career_ops"
    ev_dir = repo_root / "config" / "evidence"
    co_root.mkdir(parents=True, exist_ok=True)
    ev_dir.mkdir(parents=True, exist_ok=True)

    roles = primary_roles or ["Software Engineer", "Data Engineer", "Data Analyst"]
    locs = [loc.strip().lower() for loc in (primary_locations or [location, "remote"]) if loc.strip()]

    # 1. Write config/career_ops/config/profile.yml
    profile_yml_path = co_root / "config" / "profile.yml"
    profile_yml_path.parent.mkdir(parents=True, exist_ok=True)
    roles_yaml = "\n".join(f'    - "{r}"' for r in roles)
    profile_yml_content = f"""# Generated by `cockpit onboard`
candidate:
  full_name: "{name}"
  email: "{email}"
  phone: "{phone}"
  location: "{location}"
  linkedin: "{linkedin}"
  portfolio_url: "{portfolio}"
  github: "{github}"

target_roles:
  primary:
{roles_yaml}
  archetypes:
    - name: "Software / Backend Engineer"
      level: "Early-Career / Mid"
      fit: "primary"
    - name: "Data / Analytics Engineer"
      level: "Early-Career / Mid"
      fit: "primary"
    - name: "Product / Business Analyst"
      level: "Early-Career / Mid"
      fit: "secondary"

compensation:
  target_range: "{target_comp}"
  currency: "USD"
  minimum: "{min_comp}"
  location_flexibility: "Flexible / Remote or Hybrid"

location:
  country: "United States"
  city: "{location}"
  authorized_in: ["United States"]
  needs_sponsorship: false

language:
  output: en

spend_tier: economy
"""
    profile_yml_path.write_text(profile_yml_content, "utf-8")

    # 2. Update candidate block in config/evidence/01_master_career_history.yaml
    history_path = ev_dir / "01_master_career_history.yaml"
    if history_path.exists():
        hist_text = history_path.read_text("utf-8")
        hist_text = re.sub(r'(^\s*name:\s*)"[^"]*"', rf'\1"{name}"', hist_text, count=1, flags=re.MULTILINE)
        if email:
            hist_text = re.sub(r'(^\s*email:\s*)"[^"]*"', rf'\1"{email}"', hist_text, count=1, flags=re.MULTILINE)
        if phone:
            hist_text = re.sub(r'(^\s*phone:\s*)"[^"]*"', rf'\1"{phone}"', hist_text, count=1, flags=re.MULTILINE)
        if location:
            hist_text = re.sub(r'(^\s*location:\s*)"[^"]*"', rf'\1"{location}"', hist_text, count=1, flags=re.MULTILINE)
        history_path.write_text(hist_text, "utf-8")

    # 3. Write config/scoring_rules.local.toml for personalized outreach & location boosts
    local_rules_path = repo_root / "config" / "scoring_rules.local.toml"
    locs_toml = ", ".join(f'"{l}"' for l in locs)
    local_rules_content = f"""# Personalized overrides generated by `cockpit onboard` (gitignored)
[location]
primary_points = 10.0
primary = [{locs_toml}]

[outreach]
candidate_name = "{name}"
template = \"\"\"Hi — I'm {name}. I just applied for the {{title}} role at {{company}} and wanted to reach out directly. \\
My background is a strong fit for the work, and I'd welcome a quick 10-minute chat, \\
or a pointer to the right person to talk to.\"\"\"
"""
    local_rules_path.write_text(local_rules_content, "utf-8")

    return {
        "profile_yml": profile_yml_path,
        "history_yaml": history_path,
        "local_rules_toml": local_rules_path,
    }


def import_user_resumes(
    source_path: Path | str,
    repo_root: Path = REPO_ROOT,
) -> dict[str, Any]:
    """Imports one or more user resumes (.pdf, .docx, .md, .txt) into the workspace.

    Contract:
    - Extracts text from all discovered resumes under `source_path`.
    - Updates `config/career_ops/cv.md` and `config/evidence/base_resumes/`.
    - Copies/compiles PDF and DOCX files into `resumes/` so archetype staging works immediately.
    """
    seed_user_layer(repo_root=repo_root, overwrite=False)
    src = Path(source_path).expanduser().resolve()
    if not src.exists():
        raise FileNotFoundError(f"Resume path does not exist: {src}")

    files: list[Path] = []
    if src.is_file():
        files = [src]
    else:
        for ext in ("*.md", "*.pdf", "*.docx", "*.txt"):
            files.extend(sorted(src.glob(ext)))
        files = [f for f in files if not f.name.startswith("~$")]

    if not files:
        raise ValueError(f"No .pdf, .docx, .md, or .txt resume files found in {src}")

    co_root = repo_root / "config" / "career_ops"
    ev_dir = repo_root / "config" / "evidence"
    base_res_dir = ev_dir / "base_resumes"
    resumes_out = repo_root / "resumes"
    base_res_dir.mkdir(parents=True, exist_ok=True)
    resumes_out.mkdir(parents=True, exist_ok=True)

    extracted_docs: list[tuple[Path, str]] = []
    for f in files:
        txt = extract_resume_text(f)
        if txt:
            extracted_docs.append((f, txt))
            if f.suffix.lower() in (".pdf", ".docx"):
                dest = resumes_out / f.name
                if f.resolve() != dest.resolve():
                    shutil.copy2(f, dest)

    if not extracted_docs:
        raise ValueError(f"Could not extract text from any resume in {src}")

    # Primary resume populates cv.md
    primary_file, primary_text = extracted_docs[0]
    heuristics = parse_contact_heuristics(primary_text)
    if not primary_text.lstrip().startswith("#"):
        cv_markdown = f"# {heuristics['name']}\n\n{primary_text}\n"
    else:
        cv_markdown = f"{primary_text}\n"

    cv_path = co_root / "cv.md"
    cv_path.parent.mkdir(parents=True, exist_ok=True)
    cv_path.write_text(cv_markdown, "utf-8")

    # Populate base_resumes/*.md
    default_stems = ["swe_fullstack_ai", "data_analytics_engineering", "analytics_bi_product"]
    for idx, (orig_file, txt) in enumerate(extracted_docs):
        md_body = txt if txt.lstrip().startswith("#") else f"# {heuristics['name']}\n\n{txt}\n"
        stem = orig_file.stem
        (base_res_dir / f"{stem}.md").write_text(md_body, "utf-8")
        if idx < len(default_stems):
            (base_res_dir / f"{default_stems[idx]}.md").write_text(md_body, "utf-8")

    # If user only supplied 1 resume, mirror it to the 3 default archetype templates so all 3 archetypes compile
    if len(extracted_docs) == 1:
        for stem_alias in default_stems:
            (base_res_dir / f"{stem_alias}.md").write_text(cv_markdown, "utf-8")

    compiled = build_base_resumes(evidence_dir=ev_dir, output_dir=resumes_out)
    return {
        "imported_files": [str(f) for f, _ in extracted_docs],
        "cv_path": str(cv_path),
        "heuristics": heuristics,
        "compiled_resumes": [str(p) for p in compiled],
    }
