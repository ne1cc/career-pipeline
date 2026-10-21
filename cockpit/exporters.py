from __future__ import annotations

import html
import re
from pathlib import Path
from typing import Any

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def add_hyperlink(paragraph: Any, url: str, text: str, color: RGBColor = RGBColor(0, 102, 204), underline: bool = True) -> Any:
    """Add a clickable hyperlink to a python-docx paragraph."""
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), f"{color[0]:02x}{color[1]:02x}{color[2]:02x}")
        rPr.append(c)

    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)

    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_horizontal_border_bottom(paragraph: Any, color: str = "4F46E5", sz: str = "12") -> None:
    """Adds a stylish bottom border line under a section heading in docx."""
    pPr = paragraph._element.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)  # 1/8 pt units
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def _parse_inline_formatting(paragraph: Any, text: str, default_font_size: Pt = Pt(10), default_color: RGBColor | None = None) -> None:
    """Parses markdown links [text](url), bold **text**, and italic *text* into runs."""
    # Pattern to match [text](url), **bold**, or *italic*
    token_pattern = re.compile(r"(\[([^\]]+)\]\(([^)]+)\)|\*\*([^*]+)\*\*|\*([^*]+)\*|`([^`]+)`)")
    last_idx = 0

    for match in token_pattern.finditer(text):
        start, end = match.span()
        if start > last_idx:
            run = paragraph.add_run(text[last_idx:start])
            run.font.size = default_font_size
            if default_color:
                run.font.color.rgb = default_color

        full_token, link_text, link_url, bold_text, italic_text, code_text = match.groups()
        if link_text and link_url:
            add_hyperlink(paragraph, link_url, link_text)
        elif bold_text:
            run = paragraph.add_run(bold_text)
            run.bold = True
            run.font.size = default_font_size
            if default_color:
                run.font.color.rgb = default_color
        elif italic_text:
            run = paragraph.add_run(italic_text)
            run.italic = True
            run.font.size = default_font_size
            if default_color:
                run.font.color.rgb = default_color
        elif code_text:
            run = paragraph.add_run(code_text)
            run.font.name = "Consolas"
            run.font.size = default_font_size
            run.font.color.rgb = RGBColor(60, 60, 60)

        last_idx = end

    if last_idx < len(text):
        run = paragraph.add_run(text[last_idx:])
        run.font.size = default_font_size
        if default_color:
            run.font.color.rgb = default_color


def render_docx(markdown_text: str, output_path: Path | str, title: str = "Document") -> Path:
    """Renders a clean, ATS-compliant, single-column Microsoft Word document from Markdown."""
    out = Path(output_path).expanduser().resolve()
    out.parent.mkdir(parents=True, exist_ok=True)

    doc = docx.Document()

    # Set page margins to tight 0.5 inches for 1-page fit
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Base styling
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)  # Slate-800

    lines = markdown_text.splitlines()
    in_header = True

    for line in lines:
        raw_line = line.strip()
        if not raw_line or raw_line == "---":
            continue

        if raw_line.startswith("# ") and in_header:
            # Candidate Name / Document Title
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(raw_line[2:].strip())
            run.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(15, 23, 42)
            continue

        if in_header and ("@" in raw_line or "linkedin.com" in raw_line or "github.com" in raw_line or "(" in raw_line):
            # Contact info line
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(8)
            _parse_inline_formatting(p, raw_line, default_font_size=Pt(9.5), default_color=RGBColor(71, 85, 105))
            in_header = False
            continue

        in_header = False

        if raw_line.startswith("## "):
            # Section Heading (e.g. EDUCATION, TECHNICAL SKILLS, EXPERIENCE)
            heading_text = raw_line[3:].strip().upper()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(heading_text)
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(15, 23, 42)
            add_horizontal_border_bottom(p, color="2563EB", sz="8")
        elif raw_line.startswith("### "):
            # Sub-heading (e.g. Project title or Role)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_with_next = True
            _parse_inline_formatting(p, raw_line[4:].strip(), default_font_size=Pt(10.5))
        elif raw_line.startswith("- ") or raw_line.startswith("* "):
            # Bullet point
            bullet_text = raw_line[2:].strip()
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            _parse_inline_formatting(p, bullet_text, default_font_size=Pt(9.5))
        else:
            # Regular paragraph / institution line
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            _parse_inline_formatting(p, raw_line, default_font_size=Pt(10))

    doc.save(str(out))
    return out


def markdown_to_html(markdown_text: str, title: str = "Alex Sun - Resume") -> str:
    """Converts resume/cover letter markdown into clean semantic HTML with single-page print CSS."""
    # Process line by line into HTML
    html_lines = []
    lines = markdown_text.splitlines()
    in_list = False

    for line in lines:
        s = line.strip()
        if not s or s == "---":
            if in_list:
                html_lines.append("</ul>")
                in_list = False
            continue

        # Convert inline markdown
        # links [text](url)
        formatted = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r'<a href="\2">\1</a>', html.escape(s))
        # unescape links inside href
        formatted = re.sub(r'href="([^"]+)"', lambda m: f'href="{html.unescape(m.group(1))}"', formatted)
        # bold **text**
        formatted = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", formatted)
        # italic *text*
        formatted = re.sub(r"\*([^*]+)\*", r"<em>\1</em>", formatted)
        # code `text`
        formatted = re.sub(r"`([^`]+)`", r"<code>\1</code>", formatted)

        if s.startswith("# "):
            if in_list:
                html_lines.append("</ul>")
                in_list = False
            title_text = re.sub(r"\*\*([^*]+)\*\*", r"\1", s[2:].strip())
            html_lines.append(f'<h1 class="candidate-name">{html.escape(title_text)}</h1>')
        elif s.startswith("## "):
            if in_list:
                html_lines.append("</ul>")
                in_list = False
            heading_text = s[3:].strip().upper()
            html_lines.append(f'<h2 class="section-title">{html.escape(heading_text)}</h2>')
        elif s.startswith("### "):
            if in_list:
                html_lines.append("</ul>")
                in_list = False
            html_lines.append(f'<h3 class="item-title">{formatted[4:].strip()}</h3>')
        elif s.startswith("- ") or s.startswith("* "):
            if not in_list:
                html_lines.append('<ul class="bullet-list">')
                in_list = True
            html_lines.append(f"<li>{formatted[2:].strip()}</li>")
        else:
            if in_list:
                html_lines.append("</ul>")
                in_list = False
            if "@" in s or "linkedin.com" in s or "github.com" in s:
                html_lines.append(f'<div class="contact-line">{formatted}</div>')
            else:
                html_lines.append(f"<p>{formatted}</p>")

    if in_list:
        html_lines.append("</ul>")

    body_content = "\n".join(html_lines)

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>{html.escape(title)}</title>
<style>
  @page {{
    size: letter;
    margin: 0.42in 0.48in;
  }}
  * {{
    box-sizing: border-box;
  }}
  body {{
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
    color: #1e293b;
    background: #ffffff;
    font-size: 9.5pt;
    line-height: 1.35;
    margin: 0;
    padding: 0;
  }}
  h1.candidate-name {{
    font-size: 16pt;
    font-weight: 700;
    text-align: center;
    color: #0f172a;
    margin: 0 0 2pt 0;
    letter-spacing: -0.02em;
  }}
  .contact-line {{
    font-size: 9pt;
    text-align: center;
    color: #475569;
    margin: 0 0 8pt 0;
  }}
  .contact-line a {{
    color: #2563eb;
    text-decoration: none;
  }}
  .contact-line a:hover {{
    text-decoration: underline;
  }}
  h2.section-title {{
    font-size: 10.5pt;
    font-weight: 700;
    color: #0f172a;
    border-bottom: 1.5px solid #2563eb;
    padding-bottom: 2pt;
    margin: 8pt 0 4pt 0;
    text-transform: uppercase;
    letter-spacing: 0.05em;
  }}
  h3.item-title {{
    font-size: 9.8pt;
    font-weight: 600;
    color: #1e293b;
    margin: 5pt 0 1pt 0;
  }}
  p {{
    margin: 2pt 0;
  }}
  ul.bullet-list {{
    margin: 1pt 0 4pt 0;
    padding-left: 14pt;
  }}
  ul.bullet-list li {{
    margin: 1.5pt 0;
    line-height: 1.3;
  }}
  a {{
    color: #2563eb;
    text-decoration: none;
  }}
  strong {{
    color: #0f172a;
    font-weight: 600;
  }}
  code {{
    font-family: Consolas, Menlo, Monaco, monospace;
    font-size: 8.5pt;
    background: #f1f5f9;
    padding: 1pt 3pt;
    border-radius: 3pt;
    color: #334155;
  }}
</style>
</head>
<body>
{body_content}
</body>
</html>"""


def render_pdf(markdown_text: str, output_path: Path | str, title: str = "Document") -> Path:
    """Renders a pixel-perfect 1-page PDF using Playwright headless Chromium."""
    from playwright.sync_api import sync_playwright

    out = Path(output_path).expanduser().resolve()
    out.parent.mkdir(parents=True, exist_ok=True)

    html_content = markdown_to_html(markdown_text, title=title)

    with sync_playwright() as p:
        browser = p.chromium.launch()
        try:
            page = browser.new_page()
            page.set_content(html_content, wait_until="load")
            page.pdf(
                path=str(out),
                format="Letter",
                print_background=True,
                margin={"top": "0.42in", "bottom": "0.42in", "left": "0.48in", "right": "0.48in"},
            )
        finally:
            browser.close()

    return out
